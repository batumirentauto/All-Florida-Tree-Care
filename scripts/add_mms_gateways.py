#!/usr/bin/env python3
"""
scripts/add_mms_gateways.py

Identifies mobile carriers and generates email-to-MMS gateway addresses
for all contractors in both:
1. /Users/workpc/Documents/google-maps-scraper/output/new_scrape/
2. /Users/workpc/Documents/google-maps-scraper/output/tree_leads/

Adds columns 'Оператор связи' and 'MMS Gateway' to both companies.csv and companies.xlsx.
"""

import os
import re
import csv
import time
import json
import urllib.request
import urllib.error
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor, as_completed
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = '/Users/workpc/Documents/google-maps-scraper/output'
FOLDERS = ['new_scrape', 'tree_leads']

TOLL_FREE_PREFIXES = {'800', '888', '877', '866', '855', '844', '833'}

# In-memory cache for NPA-NXX-X lookups
PREFIX_CACHE = {}

def clean_phone_digits(phone_str):
    if not phone_str:
        return ""
    digits = re.sub(r'\D', '', phone_str)
    if len(digits) == 11 and digits.startswith('1'):
        digits = digits[1:]
    return digits

def map_telecom_to_carrier_and_gateway(company_name, ocn):
    """Maps LERG telecom entity name and OCN to major US consumer carrier brand & MMS gateway."""
    cn = (company_name or '').upper()
    
    # 1. Verizon Wireless
    if any(k in cn for k in ['VERIZON', 'CELLCO', 'ALLTEL', 'AIRTOUCH', 'PRIMECO', 'BELL ATLANTIC MOBILE', 'GTE WIRELESS']):
        return 'Verizon Wireless', '@vzwpix.com'
    
    # 2. AT&T Mobility
    if any(k in cn for k in ['NEW CINGULAR', 'CINGULAR', 'AT&T WIRELESS', 'AT&T MOBILITY', 'SBC WIRELESS', 'BELLSOUTH MOBILITY', 'SOUTHWESTERN BELL MOBILITY']):
        return 'AT&T Mobility', '@mms.att.net'
    
    # 3. T-Mobile USA
    if any(k in cn for k in ['T-MOBILE', 'METROPCS', 'VOICESTREAM', 'OMNIPOINT', 'POWERTEL', 'AERIAL', 'COOK TELECOM']):
        return 'T-Mobile', '@tmomail.net'
    
    # 4. Sprint
    if any(k in cn for k in ['SPRINT', 'NEXTEL']):
        return 'Sprint (T-Mobile)', '@pm.sprint.com'
    
    # 5. Cricket Wireless
    if 'CRICKET' in cn:
        return 'Cricket Wireless', '@mms.cricketwireless.net'
    
    # 6. Boost Mobile
    if 'BOOST' in cn:
        return 'Boost Mobile', '@myboostmobile.com'
    
    # 7. US Cellular
    if any(k in cn for k in ['UNITED STATES CELLULAR', 'US CELLULAR']):
        return 'US Cellular', '@mms.uscc.net'
    
    # 8. Spectrum Mobile (runs on Verizon)
    if 'CHARTER' in cn or 'SPECTRUM' in cn:
        return 'Spectrum Mobile', '@vzwpix.com'
    
    # 9. Xfinity Mobile / Comcast
    if 'COMCAST' in cn:
        return 'Xfinity Mobile', '@mypixmessages.com'
    
    # 10. VoIP / Virtual Carriers
    if any(k in cn for k in ['BANDWIDTH', 'TWILIO', 'PEERLESS', 'LEVEL 3', 'NEUTRAL TANDEM', 'INTELIQUENT', 'SINCH', 'TELNYX', 'VONAGE', 'RINGCENTRAL']):
        return 'VoIP / Virtual', 'Нет (VoIP)'
    
    # 11. Regional Landline Incumbents (ILEC / CLEC)
    if any(k in cn for k in ['BELLSOUTH', 'FRONTIER', 'CENTURYLINK', 'SOUTHWESTERN BELL', 'PACIFIC BELL', 'EMBARQ', 'WINDSTREAM', 'LUMEN', 'VERIZON FLORIDA LLC', 'TELEPORT']):
        short_name = company_name.split()[0] if company_name else 'Landline'
        return f'{short_name} (Landline)', 'Нет (Landline)'
    
    # Generic Wireless detection
    if 'WIRELESS' in cn or 'CELLULAR' in cn or 'PCS' in cn or 'MOBILE' in cn:
        return company_name[:25], 'Нет (Unknown Wireless)'
    
    return company_name[:25] if company_name else 'Landline', 'Нет (Landline)'

def lookup_prefix(npa, nxx, x):
    """Queries Local Calling Guide XML API for the NPA-NXX-X block."""
    cache_key = f"{npa}-{nxx}-{x}"
    if cache_key in PREFIX_CACHE:
        return PREFIX_CACHE[cache_key]

    url = f"https://localcallingguide.com/xmlprefix.php?npa={npa}&nxx={nxx}&x={x}"
    req = urllib.request.Request(url, headers={'User-Agent': 'AllFloridaTreeCare-CarrierAudit/1.0'})
    
    try:
        with urllib.request.urlopen(req, timeout=6) as resp:
            root = ET.fromstring(resp.read())
            c_elem = root.find('.//company-name')
            ocn_elem = root.find('.//ocn')
            c_name = c_elem.text if c_elem is not None else ''
            ocn = ocn_elem.text if ocn_elem is not None else ''
            res = map_telecom_to_carrier_and_gateway(c_name, ocn)
            PREFIX_CACHE[cache_key] = res
            return res
    except Exception as e:
        # Fallback query without block 'x'
        try:
            url_fallback = f"https://localcallingguide.com/xmlprefix.php?npa={npa}&nxx={nxx}"
            req_fb = urllib.request.Request(url_fallback, headers={'User-Agent': 'AllFloridaTreeCare-CarrierAudit/1.0'})
            with urllib.request.urlopen(req_fb, timeout=6) as resp:
                root = ET.fromstring(resp.read())
                c_elem = root.find('.//company-name')
                ocn_elem = root.find('.//ocn')
                c_name = c_elem.text if c_elem is not None else ''
                ocn = ocn_elem.text if ocn_elem is not None else ''
                res = map_telecom_to_carrier_and_gateway(c_name, ocn)
                PREFIX_CACHE[cache_key] = res
                return res
        except Exception:
            pass

    PREFIX_CACHE[cache_key] = ('Unknown', 'Нет (Lookup Error)')
    return PREFIX_CACHE[cache_key]

def resolve_phone_mms(phone_raw):
    """Takes raw phone string and returns (Carrier, MMS_Gateway)."""
    digits = clean_phone_digits(phone_raw)
    if not digits or len(digits) != 10:
        return 'Нет номера', 'Нет номера'
    
    area_code = digits[:3]
    if area_code in TOLL_FREE_PREFIXES:
        return 'Toll-free (800/833)', 'Нет (Toll-free)'
    
    npa = digits[:3]
    nxx = digits[3:6]
    x = digits[6]
    
    carrier, domain = lookup_prefix(npa, nxx, x)
    if domain.startswith('@'):
        mms_gateway = f"{digits}{domain}"
    else:
        mms_gateway = domain
        
    return carrier, mms_gateway

def process_folder(folder_name):
    folder_path = os.path.join(BASE_DIR, folder_name)
    csv_path = os.path.join(folder_path, 'companies.csv')
    xlsx_path = os.path.join(folder_path, 'companies.xlsx')

    if not os.path.exists(csv_path):
        print(f"Skipping {folder_name}: {csv_path} not found.")
        return

    print(f"\n=======================================================")
    print(f"Processing folder: {folder_name}")
    print(f"=======================================================")

    with open(csv_path, 'r', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        header = next(reader)
        rows = list(reader)

    phone_idx = header.index('Контактный номер')
    
    # Determine insertion position
    if 'Наличие iMessage' in header:
        insert_pos = header.index('Наличие iMessage') + 1
    else:
        insert_pos = phone_idx + 1

    # Check if already present
    carrier_already = 'Оператор связи' in header
    mms_already = 'MMS Gateway' in header

    if carrier_already and mms_already:
        print("Columns 'Оператор связи' and 'MMS Gateway' already exist. Updating values...")
        c_idx = header.index('Оператор связи')
        m_idx = header.index('MMS Gateway')
        new_header = header
        new_rows = rows
    else:
        new_header = header[:insert_pos] + ['Оператор связи', 'MMS Gateway'] + header[insert_pos:]
        c_idx = insert_pos
        m_idx = insert_pos + 1
        new_rows = []
        for r in rows:
            new_r = r[:insert_pos] + ['', ''] + r[insert_pos:]
            new_rows.append(new_r)

    print(f"Resolving carrier & MMS gateways for {len(new_rows)} contractors...")

    # Collect phone numbers for concurrent resolution
    tasks = []
    with ThreadPoolExecutor(max_workers=5) as executor:
        future_to_idx = {
            executor.submit(resolve_phone_mms, r[phone_idx]): i
            for i, r in enumerate(new_rows)
        }
        for future in as_completed(future_to_idx):
            i = future_to_idx[future]
            try:
                carrier, mms_gw = future.result()
                new_rows[i][c_idx] = carrier
                new_rows[i][m_idx] = mms_gw
            except Exception as e:
                new_rows[i][c_idx] = 'Error'
                new_rows[i][m_idx] = 'Нет'

    # Save CSV
    print(f"Saving updated CSV: {csv_path}")
    with open(csv_path, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(new_header)
        writer.writerows(new_rows)

    # Save styled XLSX
    print(f"Saving styled XLSX: {xlsx_path}")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Подрядчики"

    header_fill = PatternFill(start_color="1B4332", end_color="1B4332", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    ws.append(new_header)
    for col_idx in range(1, len(new_header) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    mms_col_1based = m_idx + 1
    carrier_col_1based = c_idx + 1

    for row_idx, r in enumerate(new_rows, start=2):
        ws.append(r)
        ws.row_dimensions[row_idx].height = 110
        for col_idx, val in enumerate(r, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = cell_font
            cell.border = thin_border
            
            # Highlight active MMS Gateways
            if col_idx == mms_col_1based:
                cell.alignment = Alignment(horizontal="center", vertical="top")
                if "@" in str(val):
                    cell.font = Font(name="Arial", size=10, bold=True, color="1B5E20")
                    cell.fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
                else:
                    cell.font = Font(name="Arial", size=10, color="777777")
            elif col_idx == carrier_col_1based:
                cell.alignment = Alignment(horizontal="center", vertical="top")
                cell.font = Font(name="Arial", size=10, bold=True if "@" in str(r[m_idx]) else False)
            elif 'imessage' in new_header[col_idx - 1].lower():
                cell.alignment = Alignment(horizontal="center", vertical="top")
                if val == "Да":
                    cell.font = Font(name="Arial", size=10, bold=True, color="1B5E20")
                    cell.fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
            elif col_idx == len(new_header):  # Сообщение
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif col_idx in [3, 4, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top")

    # Column widths
    widths = {
        'Название': 32,
        'Ссылка на Google Maps': 25,
        'Оценка': 10,
        'Количество отзывов': 12,
        'Наличие сайта': 12,
        'Ссылка на сайт': 28,
        'Проблема с сайтом': 20,
        'Контактный номер': 18,
        'Наличие iMessage': 18,
        'Оператор связи': 22,
        'MMS Gateway': 28,
        'Адрес': 35,
        'Штат': 10,
        'Сообщение для рассылки': 75
    }
    for col_idx, h in enumerate(new_header, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(h, 20)

    ws.freeze_panes = 'B2'
    wb.save(xlsx_path)

    # Statistics
    mms_count = sum(1 for r in new_rows if "@" in str(r[m_idx]))
    print(f"Done for {folder_name}: {mms_count}/{len(new_rows)} active MMS gateway addresses identified.")

def main():
    for f in FOLDERS:
        process_folder(f)
    print("\nAll folders processed successfully!")

if __name__ == '__main__':
    main()
