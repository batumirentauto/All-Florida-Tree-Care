#!/usr/bin/env python3
"""
scripts/add_outreach_messages.py
Adds personalized outreach messages as a new column to companies.csv and companies.xlsx
in /Users/workpc/Documents/google-maps-scraper/output/new_scrape/.
Includes 'Reply STOP to opt out' at the end of every message.
"""

import os
import csv
import openpyxl

MESSAGES = {
    'A Integrity Tree Surgeons Inc': (
        "Hey Steve, saw your 4.7 rating in Jacksonville and how much locals value that your Vermeer track grinder tackles massive hardwood stumps next to foundations without tearing up lawns.\n\n"
        "Put together a fast mobile site concept for A Integrity Tree Surgeons featuring your commercial equipment (Vermeer SC362, heavy rigging gear), spotless yard cleanups, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Tally Trees': (
        "Hey Tally Trees team, saw your 5.0 rating across Tallahassee and how much homeowners respect your hydraulic mobile crane work — safely hoisting colossal pine tops right over tight residential rooflines.\n\n"
        "Put together a fast mobile site concept for Tally Trees featuring your heavy crane rigging, CAT loader on ground mats, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'I NOAH GUY TREE & OUTDOORS LLC': (
        "Hey Noah, saw your 5.0 rating in Jacksonville and how much locals appreciate your quick response for hazardous takedowns and that your crew leaves properties cleaner than when you arrived.\n\n"
        "Put together a fast mobile site concept for I NOAH GUY Tree & Outdoors featuring your aerial bucket rig, commercial chipping trailer, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Hawkinberry Tree Company': (
        "Hey Jeff, saw your 5.0 rating across Jacksonville Beach and how much clients value that your rubber-tracked spider lift squeezes through narrow 36-inch gates to reach canopy branches over utility wires with zero turf impact.\n\n"
        "Put together a fast mobile site concept for Hawkinberry Tree Company featuring your Platform Basket spider lift, tight-access canopy pruning, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Family Tree Service LLC': (
        "Hey Jacob, saw your 4.6 rating across Camden County and how much homeowners rely on your family crew for emergency storm takedowns and thorough whole-tree chipping without surprise fees.\n\n"
        "Put together a fast mobile site concept for Family Tree Service featuring your heavy-duty chipper, aerial boom access, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'G.R.O.W. Outdoor Services, Inc': (
        "Hey Donnie, saw your 4.9 rating in Jacksonville and how much locals highlight your crew's punctuality, surgical trimming over fence lines, and leaving zero debris behind on driveways.\n\n"
        "Put together a fast mobile site concept for G.R.O.W. Outdoor Services featuring your commercial climbing and rigging fleet, lawn preservation standards, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Bramlitt Tree Service LLC': (
        "Hey Derek, saw your 5.0 rating in Clay County and how much locals admire your signature diamond-cut Sylvester palm shaving and controlled sectional pruning away from residential rooflines.\n\n"
        "Put together a fast mobile site concept for Bramlitt Tree Service featuring your artisan palm restoration, hazardous tree surgery, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Allwood Tree Service': (
        "Hey James, saw your 4.9 rating across St. Johns & Jacksonville and how much property owners value your precision live oak felling and Vermeer stump grinding without tearing up sod or landscape curbing.\n\n"
        "Put together a fast mobile site concept for Allwood Tree Service featuring your Vermeer milling equipment, residential oak pruning, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'S & S Tree Service': (
        "Hey Scott, saw your 4.7 rating in Jacksonville and how much locals value your crew's heavy rigging capability to take down massive rotted trees safely without damage to neighboring properties.\n\n"
        "Put together a fast mobile site concept for S & S Tree Service featuring your heavy canopy rigging setup, direct owner dispatch, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'At Ease Land & Tree Services LLC': (
        "Hey John, saw your 4.9 rating in Cairo & Thomasville and how much locals appreciate your prompt communication, fair written quotes, and efficient storm damage clearing.\n\n"
        "Put together a fast mobile site concept for At Ease Land & Tree Services featuring your commercial felling gear, below-grade stump removal, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Panhandle Tree Pros LLC': (
        "Hey Mike, saw your 5.0 rating across the Panhandle and how much locals respect your fearless canopy climbing and controlled sectional takedowns around tight rural cabins and power lines.\n\n"
        "Put together a fast mobile site concept for Panhandle Tree Pros featuring your high-reach aerial equipment, brush chipping rig, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Limbwalkers Tree & Crane Service': (
        "Hey Caleb, saw your 4.9 rating in Burke County and how much clients value having your dedicated crane service for lifting hazardous mountain hardwoods safely away from steep residential rooflines.\n\n"
        "Put together a fast mobile site concept for Limbwalkers Tree & Crane featuring your hydraulic crane picks, tracked chipper fleet, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Always Grinding Tree & Stump Removal': (
        "Hey Dustin, saw your 5.0 rating around Lake Lanier and how much homeowners appreciate that your commercial stump grinder pulverizes deep root balls below grade with total clean backfilling.\n\n"
        "Put together a fast mobile site concept for Always Grinding featuring your high-horsepower track grinder, precision hazard removals, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'T & D Professional Tree Removal': (
        "Hey Timothy, saw your 4.9 rating along the Mississippi Gulf Coast and how much locals rely on your crew for urgent storm clearing and honest, upfront pricing after heavy weather.\n\n"
        "Put together a fast mobile site concept for T & D Professional Tree Removal featuring your heavy timber felling gear, 24/7 emergency dispatch, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Broken Branch Tree Removal': (
        "Hey Fred, saw your 5.0 rating across Fulton County and how much homeowners value your technical rope rigging when dismantling split pine and oak trees hanging dangerously over roofs.\n\n"
        "Put together a fast mobile site concept for Broken Branch Tree Removal featuring your controlled canopy lowering gear, fast cleanup crew, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Superior Tree Service aka little tree lady': (
        "Hey Naomi, saw your 4.9 rating in Senoia & Coweta County and how much locals praise your detailed pruning eye, honest advice on tree health, and meticulous yard cleanup.\n\n"
        "Put together a fast mobile site concept for Superior Tree Service featuring your canopy preservation methods, hazardous limb clearance, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'US TREE REMOVAL': (
        "Hey Cass, saw your 5.0 rating in Dearborn & Wayne County and how much clients respect that your team safely dismantles huge diseased trees in tight residential lots without touching fences.\n\n"
        "Put together a fast mobile site concept for US Tree Removal featuring your heavy rigging truck, rapid wood hauling, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Columbus Georgia Tree Removal': (
        "Hey Russell, saw your 5.0 rating in Columbus and how much homeowners value your quick response time for hazardous leaning trees and clear, transparent estimates without hidden charges.\n\n"
        "Put together a fast mobile site concept for Columbus Georgia Tree Removal featuring your aerial bucket access, high-capacity chipper, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Eaton Tree Removal': (
        "Hey Chris, saw your 4.9 rating in Blue Ridge and how much mountain cabin owners rely on your crew for precision takedowns on steep grades and clearing dead pines before winter storms.\n\n"
        "Put together a fast mobile site concept for Eaton Tree Removal featuring your slope-access rigging gear, comprehensive site cleanup, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Maldonado Tree Service': (
        "Hey Mr. Maldonado, saw your 4.7 rating across Houston and how much homeowners depend on your crew for heavy post-storm limb clearing and precision oak pruning away from roofs and power drops.\n\n"
        "Put together a fast mobile site concept for Maldonado Tree Service featuring your heavy bucket truck, high-volume brush chipping, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Tafoya\'s Tree Service & Stump Grinding': (
        "Hey Nester, saw your 4.8 rating in Houston and how much locals value your prompt communication, fair pricing, and clean stump grinding that leaves yards ready for new grass.\n\n"
        "Put together a fast mobile site concept for Tafoya's Tree Service featuring your commercial stump grinding machinery, controlled takedowns, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'L&M Tree Service': (
        "Hey Logan, saw your 4.7 rating in Muskogee and how much property owners appreciate your crew's fast turnaround on dangerous dead trees and leaving the work area completely raked clean.\n\n"
        "Put together a fast mobile site concept for L&M Tree Service featuring your bucket truck capabilities, commercial wood hauling, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'M Rosales Tree Service': (
        "Hey Rodrigo, saw your 5.0 rating across Houston and how much clients respect your punctuality, careful branch lowering over fences, and competitive written bids.\n\n"
        "Put together a fast mobile site concept for M Rosales Tree Service featuring your controlled rigging setup, complete debris haul-away, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Sutton\'s Tree Service Inc': (
        "Hey Sutton's Tree Service team, saw your 4.8 rating across Onslow County and how much military families and locals value your honesty, fast storm dispatch, and thorough cleanups.\n\n"
        "Put together a fast mobile site concept for Sutton's Tree Service featuring your aerial takedown gear, deep stump grinding, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Allen\'s Tree Service': (
        "Hey Allen's Tree Service team, saw your 5.0 rating in Youngsville & Franklin County and how much homeowners praise your prompt estimates, careful property protection, and spotless raking.\n\n"
        "Put together a fast mobile site concept for Allen's Tree Service featuring your precision canopy pruning, hazardous takedown rig, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'WoodFellas Tree Service': (
        "Hey John, saw your 5.0 rating across the Tri-Cities and how much homeowners respect your crew's skill dismantling tall hillside trees without dinging gutters or landscaping.\n\n"
        "Put together a fast mobile site concept for WoodFellas Tree Service featuring your technical rigging gear, complete debris removal, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Stone\'s Tree Service and Stump Grinding': (
        "Hey Christian, saw your 4.9 rating in Greenville County and how much locals value that you handle massive multi-tree clearing jobs quickly with zero lawn destruction.\n\n"
        "Put together a fast mobile site concept for Stone's Tree Service featuring your heavy equipment fleet, below-grade stump grinding, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'JSI Tree Service LLC': (
        "Hey Nathan, saw your 4.9 rating in Morganton and how much clients appreciate your fair pricing on tricky tree removals and leaving yards free of twigs and ruts.\n\n"
        "Put together a fast mobile site concept for JSI Tree Service featuring your residential takedown setup, commercial brush chipping, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Pro Beaver Tree Service': (
        "Hey Ivan, saw your 5.0 rating in Waxhaw & Union County and how much homeowners praise your exact directional felling, polite crew, and leaving driveways blown completely clean.\n\n"
        "Put together a fast mobile site concept for Pro Beaver Tree Service featuring your precision climbing rig, stump grinding unit, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'TLG Tree Service LLC': (
        "Hey Oliver, saw your 4.9 rating across the Triangle and how much locals value your rapid turnaround for storm hazards and careful crown thinning over suburban houses.\n\n"
        "Put together a fast mobile site concept for TLG Tree Service featuring your aerial bucket access, high-capacity debris trailer, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Furstie Tree & Stump, LLC': (
        "Hey Nate, saw your 5.0 rating in Orange County and how much clients appreciate your specialized track stump grinding that erases root flares below grade without tearing up established turf.\n\n"
        "Put together a fast mobile site concept for Furstie Tree & Stump featuring your commercial grinding equipment, hazardous tree removal, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Bartolome\'s Tree service': (
        "Hey Bartolome, saw your 4.8 rating in Sanford and how much local property owners praise your hard work on tough multi-acre clearances and honest, upfront quotes.\n\n"
        "Put together a fast mobile site concept for Bartolome's Tree Service featuring your heavy felling gear, complete wood haul-away, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Wateree Tree RemovalLLC': (
        "Hey David, saw your work across Kershaw County and how much locals count on your quick response for emergency storm blowdowns and straightforward, reliable service.\n\n"
        "Put together a fast mobile site concept for Wateree Tree Removal featuring your heavy rigging truck, rapid site clearance, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Tree Wisemen Upstate': (
        "Hey Evan, saw your 4.9 rating in Anderson and how much Upstate homeowners value your skilled climbing arborists who take down massive dead pines in tight spaces with surgical precision.\n\n"
        "Put together a fast mobile site concept for Tree Wisemen Upstate featuring your technical rigging gear, commercial brush chipper, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Norvell\'s Tree & Landscape Services LLC': (
        "Hey Stevie, saw your 4.8 rating in Spartanburg and how much clients appreciate your crew's speed, honest pricing, and making sure every branch and twig is cleaned up before you leave.\n\n"
        "Put together a fast mobile site concept for Norvell's Tree & Landscape featuring your bucket truck access, whole-tree hauling, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Morgan\'s Tree & Stump Removal': (
        "Hey Mike, saw your 5.0 rating in Rock Hill and how much homeowners value your precision takedowns over pools and sheds, plus grinding stumps below grade so the yard looks brand new.\n\n"
        "Put together a fast mobile site concept for Morgan's Tree & Stump Removal featuring your high-reach aerial equipment, track grinder, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Clean Cut Tree Service': (
        "Hey Wade, saw your 5.0 rating across Pierce County and how much Pacific Northwest homeowners trust your crew to dismantle giant decaying Douglas firs over tight residential fences safely.\n\n"
        "Put together a fast mobile site concept for Clean Cut Tree Service featuring your heavy rigging fleet, commercial chipping unit, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Bigfoot tree care': (
        "Hey Dave, saw your 4.9 rating around Kent & South King County and how much clients value your arborist expertise, careful limb lowering around power lines, and spotless cleanups.\n\n"
        "Put together a fast mobile site concept for Bigfoot Tree Care featuring your technical climbing gear, hazard assessment standards, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Best For Less Tree Care': (
        "Hey James, saw your 5.0 rating in Vancouver & Clark County and how much homeowners appreciate that you deliver top-tier tree removals and pruning at honest, reasonable prices without cut corners.\n\n"
        "Put together a fast mobile site concept for Best For Less Tree Care featuring your bucket boom access, high-capacity wood chipper, and a built-in 10-second price estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Affordable Tree and Palm Tree Services': (
        "Hey Francisco, saw your 4.7 rating in Sacramento and how much locals rely on your crew for skilled palm skinning, large hardwood pruning, and leaving properties totally clean.\n\n"
        "Put together a fast mobile site concept for Affordable Tree and Palm Tree Services featuring your aerial lift equipment, heavy debris hauling, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Dane’s Tree & Excavation Service': (
        "Hey Dane, saw your 5.0 rating across Marion County and how much locals respect your precision directional felling — steering multi-ton Douglas firs right between carports and campers without a scratch.\n\n"
        "Put together a fast mobile site concept for Dane's Tree & Excavation featuring your fuel-injected STIHL falling saws, high-spur climbing rig, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Patrick Myers Tree Service LLC': (
        "Hey Patrick, saw your 4.7 rating in Coos Bay and how much locals value your bucket truck reach for taking down dead coastal cedars and clearing heavy limbs away from apartment complexes and driveways.\n\n"
        "Put together a fast mobile site concept for Patrick Myers Tree Service featuring your aerial bucket boom, commercial chipping unit, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    ),
    'Holcomb Timber & Tree Services LLC': (
        "Hey Tony, saw your 5.0 rating across Lane County and how much property owners respect your logger-grade high-climbing skills when tackling 100+ ft firs and coastal redwoods in tight quarters.\n\n"
        "Put together a fast mobile site concept for Holcomb Timber & Tree featuring your technical timber rigging, heavy-duty chipper fleet, and a built-in 10-second tree cost estimator for homeowners (preview attached).\n\n"
        "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George\n\n"
        "Reply STOP to opt out"
    )
}

def get_message_for_company(cname):
    cname = cname.strip()
    if cname in MESSAGES:
        return MESSAGES[cname]
    # check alternate quotes
    cname_alt = cname.replace('’', "'")
    for k, v in MESSAGES.items():
        if k.replace('’', "'") == cname_alt:
            return v
    raise ValueError(f"No message found for company: {cname}")

def update_csv():
    csv_path = '/Users/workpc/Documents/google-maps-scraper/output/new_scrape/companies.csv'
    with open(csv_path, mode='r', encoding='utf-8') as f:
        reader = csv.reader(f)
        header = next(reader)
        rows = list(reader)

    # Check if column already present
    col_name = 'Сообщение для рассылки'
    if col_name in header:
        col_idx = header.index(col_name)
    else:
        header.append(col_name)
        col_idx = len(header) - 1

    new_rows = []
    for row in rows:
        cname = row[0].strip()
        msg = get_message_for_company(cname)
        if len(row) <= col_idx:
            row.extend([''] * (col_idx - len(row) + 1))
        row[col_idx] = msg
        new_rows.append(row)

    with open(csv_path, mode='w', encoding='utf-8', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(new_rows)
    print(f"Updated CSV: {csv_path} ({len(new_rows)} rows updated)")

def update_xlsx():
    xlsx_path = '/Users/workpc/Documents/google-maps-scraper/output/new_scrape/companies.xlsx'
    wb = openpyxl.load_workbook(xlsx_path)
    sheet = wb.active

    col_name = 'Сообщение для рассылки'
    header = [cell.value for cell in sheet[1]]
    if col_name in header:
        col_idx = header.index(col_name) + 1
    else:
        col_idx = len(header) + 1
        sheet.cell(row=1, column=col_idx, value=col_name)

    # Format header cell like others
    sheet.cell(row=1, column=col_idx).font = openpyxl.styles.Font(bold=True)

    updated_count = 0
    for r in range(2, sheet.max_row + 1):
        cname = sheet.cell(row=r, column=1).value
        if not cname:
            continue
        msg = get_message_for_company(str(cname).strip())
        cell = sheet.cell(row=r, column=col_idx, value=msg)
        cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
        updated_count += 1

    # Adjust column width
    sheet.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 65

    wb.save(xlsx_path)
    print(f"Updated XLSX: {xlsx_path} ({updated_count} rows updated)")

if __name__ == '__main__':
    update_csv()
    update_xlsx()
