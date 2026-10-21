#!/usr/bin/env python3
"""
scripts/update_user_custom_table.py

Applies showcase terminology, removes STOP opt-out, and adds Two-Step Outreach
messages to the user's custom file:
  "companies tree removers 250 - companies.csv.csv"
in /Users/workpc/Documents/google-maps-scraper/output/tree_leads/

Also outputs:
  - "companies tree removers 250 - companies.csv" (clean extension copy)
  - "companies tree removers 250 - companies.xlsx" (styled Excel workbook)
"""

import os
import re
import csv
import json
import glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

WORKSPACE = "/Users/workpc/Documents/All Florida Tree Care"
TREE_LEADS_DIR = "/Users/workpc/Documents/google-maps-scraper/output/tree_leads"
TARGET_CSV_NAME = "companies tree removers 250 - companies.csv.csv"
CLEAN_CSV_NAME = "companies tree removers 250 - companies.csv"
TARGET_XLSX_NAME = "companies tree removers 250 - companies.xlsx"
BASE_SITE_URL = "https://all-florida-tree-care-production.up.railway.app"

US_STATE_CODES = {
    'AL', 'AK', 'AZ', 'AR', 'CA', 'CO', 'CT', 'DE', 'FL', 'GA',
    'HI', 'ID', 'IL', 'IN', 'IA', 'KS', 'KY', 'LA', 'ME', 'MD',
    'MA', 'MI', 'MN', 'MS', 'MO', 'MT', 'NE', 'NV', 'NH', 'NJ',
    'NM', 'NY', 'NC', 'ND', 'OH', 'OK', 'OR', 'PA', 'RI', 'SC',
    'SD', 'TN', 'TX', 'UT', 'VT', 'VA', 'WA', 'WV', 'WI', 'WY'
}

NON_HUMAN_NAMES = {
    'houston', 'gulf', 'twisted', 'acadia', 'green', 'arbor', 'timberland', 
    'archon', 'treekeeper', 'automated', 'rainbow', 'kats', 'cutchins', 'asay',
    'campos', 'henriquez', 'lefler', 'dallas', 'austin', 'miami', 'tampa',
    'atlanta', 'charlotte', 'seattle', 'orlando', 'lead', 'lead arborist',
    'dispatcher', 'dispatch', 'owner', 'tree', 'n/a', 'unknown', 'none', 'team',
    'service', 'services', 'contractor', 'landscaping', 'lawn', 'outdoor', 'outdoors',
    'care', 'solutions', 'affordable', 'all', 'top', 'pure', 'pro', 'mr', 'dr',
    'champion', 'superior', 'precision', 'premier', 'elite', 'budget', 'local',
    'southern', 'american', 'true', 'ultimate', 'custom', 'classic', 'first',
    'national', 'metro', 'city', 'state', 'county', 'commercial', 'residential'
}

MANUAL_LOCATION_OVERRIDES = {
    'allweathertreeandlandscaping': ('South Florida', 'FL'),
    'allweathertreeandlandscapingserviceinc': ('South Florida', 'FL'),
    'cheapeguytreeandstumpremoval': ('Central Florida', 'FL'),
    'cheapguytreeandstumpremoval': ('Central Florida', 'FL'),
    'continentalcoredrillingtreeservice': ('Miami', 'FL'),
    'progreenlandscapesolutions': ('North Florida', 'FL'),
    'streeservicellc': ('South Florida', 'FL'),
    'thelawnguy': ('Central Florida', 'FL'),
    'topcuttreeservice': ('South Florida', 'FL'),
    'topshelftreeservicelandscaping': ('Central Florida', 'FL'),
    'troncosland': ('Miami', 'FL'),
    'trueclimblogistics': ('South Florida', 'FL'),
    'truetreeservice': ('Miami', 'FL'),
    'tuckerenterpriseservices': ('North Florida', 'FL'),
    'yourtreeguy': ('Central Florida', 'FL'),
}

def clean_key(name):
    if not name: return ""
    return re.sub(r'[^a-zA-Z0-9]', '', name.lower())

def clean_biz_name(raw_name):
    clean = re.sub(r'\s*,?\s*\b(LLC|Inc\.?|Corp\.?|Co\.?|L\.L\.C\.)\b', '', raw_name, flags=re.IGNORECASE)
    clean = re.sub(r'\s*-\s*Tree Service.*$', '', clean, flags=re.IGNORECASE)
    clean = re.sub(r'\s+', ' ', clean).strip()
    return clean

def load_site_configs():
    configs = {}
    pattern = os.path.join(WORKSPACE, "src/siteConfig.*.en.js")
    for filepath in glob.glob(pattern):
        slug = os.path.basename(filepath).replace("siteConfig.", "").replace(".en.js", "")
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                content = f.read()
            name_m = re.search(r'["\']?(?:businessName|name)["\']?\s*:\s*["\']([^"\']+)["\']', content)
            city_m = re.search(r'["\']?city["\']?\s*:\s*["\']([^"\']+)["\']', content)
            owner_m = re.search(r'["\']?(?:ownerFirstName|ownerName)["\']?\s*:\s*["\']([^"\']+)["\']', content)
            lead_arb_m = re.search(r'["\']?leadArborist["\']?\s*:\s*["\']([^"\']+)["\']', content)

            cname = name_m.group(1).strip() if name_m else ""
            city = city_m.group(1).strip() if city_m else ""
            owner = owner_m.group(1).strip() if owner_m else ""
            if not owner and lead_arb_m:
                owner = lead_arb_m.group(1).strip()

            item = {
                'slug': slug,
                'name': cname,
                'city': city,
                'ownerFirstName': owner
            }
            if cname:
                configs[clean_key(cname)] = item
            configs[clean_key(slug)] = item
        except Exception:
            pass
    return configs

def load_leads_data():
    leads = {}
    pattern = os.path.join(TREE_LEADS_DIR, "*/company.json")
    for filepath in glob.glob(pattern):
        folder = os.path.dirname(filepath)
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                data = json.load(f)
            data['_folder'] = folder
            cname = data.get('companyName') or data.get('name') or ""
            if cname:
                leads[clean_key(cname)] = data
            folder_name = os.path.basename(folder)
            leads[clean_key(folder_name)] = data
        except Exception:
            pass
    return leads

def find_company_folder(base_dir, company_name):
    ck = clean_key(company_name)
    if not os.path.exists(base_dir):
        return None
    for entry in os.scandir(base_dir):
        if entry.is_dir():
            if clean_key(entry.name) == ck:
                return entry.path
    for entry in os.scandir(base_dir):
        if entry.is_dir():
            eck = clean_key(entry.name)
            if ck and (ck in eck or eck in ck):
                return entry.path
    return None

def analyze_reviews_and_equipment(folder_path, config_item):
    reviews_text = ""
    reviews_path = os.path.join(folder_path, 'reviews.json') if folder_path else None
    if reviews_path and os.path.exists(reviews_path):
        try:
            revs = json.load(open(reviews_path, encoding='utf-8'))
            reviews_text = " ".join([r.get('text', '') for r in revs if isinstance(r, dict)]).lower()
        except Exception:
            pass

    has_crane = 'crane' in reviews_text
    has_spider = 'spider' in reviews_text or 'narrow' in reviews_text or 'gate' in reviews_text
    has_stump = 'stump' in reviews_text or 'grind' in reviews_text
    has_clean = any(w in reviews_text for w in ['clean', 'spotless', 'rake', 'debris', 'leaf', 'yard'])
    has_storm = any(w in reviews_text for w in ['storm', 'emergency', 'hurricane', 'fallen', 'urgent'])
    has_tight = any(w in reviews_text for w in ['tight', 'fence', 'roof', 'power', 'close', 'lean', 'hazard'])
    has_palm = 'palm' in reviews_text

    # Proof
    if has_crane:
        proof = "how safely your crew coordinates hydraulic crane picks to hoist massive canopy branches right over tight rooflines"
        fleet = "heavy crane rigging capabilities"
    elif has_spider:
        proof = "that your compact aerial lift maneuvers through tight residential gates without disturbing lawn turf"
        fleet = "compact spider lift access"
    elif has_storm:
        proof = "how rapidly your emergency crew mobilizes to clear storm hazards and hung-up limbs safely"
        fleet = "24/7 emergency dispatch and heavy timber rigging"
    elif has_tight and has_clean:
        proof = "how smoothly your crew executes precision takedowns in tight backyard spaces and leaves properties spotless"
        fleet = "controlled sectional lowering gear"
    elif has_stump:
        proof = "that your commercial stump milling grinds deep root balls below grade with total clean backfilling"
        fleet = "high-horsepower stump grinder and chipping fleet"
    elif has_palm:
        proof = "your crew's precision palm trimming and hazardous frond clearing with zero damage to underlying landscape"
        fleet = "specialized palm and aerial bucket fleet"
    else:
        proof = "how smoothly your team handles hazardous tree removals and leaves residential yards immaculately cleaned"
        fleet = "commercial rigging equipment and chipping fleet"

    # Benefit
    if has_clean:
        benefit = "spotless yard cleanups"
    elif has_stump:
        benefit = "subsurface stump milling"
    elif has_storm:
        benefit = "24/7 storm damage dispatch"
    else:
        benefit = "zero-turf-damage lawn protection"

    return proof, fleet, benefit

def extract_city(raw_name, addr, lead_item, config_item):
    ck = clean_key(raw_name)
    if ck in MANUAL_LOCATION_OVERRIDES:
        return MANUAL_LOCATION_OVERRIDES[ck][0]
    
    if addr and addr.strip() != 'Не указан':
        m = re.search(r',\s*([^,]+),\s*([A-Z]{2})\s+(\d{5})', addr)
        if m and m.group(2) in US_STATE_CODES:
            return m.group(1).strip()
        m_st = re.search(r'\b([A-Z]{2})\b\s+\d{5}', addr)
        if m_st and m_st.group(1) in US_STATE_CODES:
            parts = addr.split(',')
            city = parts[-2].strip() if len(parts) >= 2 else ''
            city = re.sub(r'\b[A-Z]{2}\b|\d{5}.*', '', city).strip()
            if city: return city

    if lead_item and lead_item.get('city') and lead_item['city'] != 'Local':
        return lead_item['city']
    if config_item and config_item.get('city') and config_item['city'] != 'Local':
        return config_item['city']

    return 'Florida'

def extract_owner(raw_name, lead_item, config_item):
    cand = None
    if config_item and config_item.get('ownerFirstName'):
        cand = config_item['ownerFirstName'].strip()
    elif lead_item and lead_item.get('ownerEn'):
        cand = lead_item['ownerEn'].strip()

    if cand:
        cand_lower = cand.lower()
        if cand_lower not in NON_HUMAN_NAMES and len(cand) > 1 and not cand.isdigit():
            return cand
    return None

def resolve_page_url(cname, lead_item, config_item):
    slug = None
    if lead_item and lead_item.get('slug'):
        slug = lead_item['slug']
    if not slug and config_item and config_item.get('slug'):
        slug = config_item['slug']

    if slug:
        return f"{BASE_SITE_URL}/{slug}"
    return "—"

def build_original_message(raw_name, rating, city, owner_name, proof, fleet, benefit):
    clean_biz = clean_biz_name(raw_name)
    try:
        r_val = float(rating)
        r_str = f"{r_val:.1f}"
    except:
        r_str = "top-rated"
        
    rating_phrase = f"your {r_str} rating" if r_str != "top-rated" else "your stellar reputation"
    location_phrase = f"in {city}" if city and city != "Florida" else "across the area"
    
    greeting = f"Hey {owner_name}," if owner_name else f"Hey {clean_biz} team,"

    p1 = f"{greeting} saw {rating_phrase} {location_phrase} and how much locals value {proof}."
    p2 = f"Put together a fast mobile showcase concept for {clean_biz} featuring your {fleet}, {benefit}, and a built-in 10-second price estimator for homeowners (preview attached)."
    p3 = "Zero pressure or cost to take a look. If you'd like to check out the full live demo on your phone, just reply 'yes' and I’ll send over the link! — George"

    return f"{p1}\n\n{p2}\n\n{p3}"

def build_4_messages(raw_name, rating, city, owner_name, fleet):
    clean_biz = clean_biz_name(raw_name)
    try:
        r_val = float(rating)
        r_str = f"{r_val:.1f}"
    except:
        r_str = "top-rated"

    rating_phrase = f"your {r_str}-star Google reviews" if r_str != "top-rated" else "your top-rated Google reviews"
    loc_phrase = f"in {city}" if city and city != "Florida" else "across the area"
    loc_city = city if city and city != "Florida" else "your local area"

    if owner_name:
        g1 = f"Hey {owner_name},"
        g2 = f"Hi {owner_name},"
        g3 = f"Good day {owner_name},"
        g4 = f"Hello {owner_name},"
    else:
        g1 = f"Hey {clean_biz} team,"
        g2 = f"Hi {clean_biz} crew,"
        g3 = f"Hello {clean_biz},"
        g4 = f"Greetings {clean_biz} team,"

    # OPTION 1: Complete Custom Showcase (Готовая витрина)
    msg1 = (
        f"{g1}\n\n"
        f"Put together a complete, custom showcase for {clean_biz}. I pulled in your real equipment photos, "
        f"verified service areas {loc_phrase}, and {rating_phrase}. "
        f"The showcase is fully responsive — loads blazing fast and looks sharp whether homeowners visit on a smartphone or a desktop computer.\n\n"
        f"Want me to send over the link to check it out? — George"
    )

    # OPTION 2: Conversion & Instant Estimator (Конверсия и смета)
    msg2 = (
        f"{g2}\n\n"
        f"Saw your tree service profile {loc_phrase} — you have a stellar reputation, but high-ticket commercial and residential jobs "
        f"need a modern showcase to close fast. Built a full interactive showcase for {clean_biz} featuring a 10-second price estimator, "
        f"your heavy equipment specs, and instant quote booking on any device.\n\n"
        f"Can I send the link over to take a quick look? Zero cost or obligation. — George"
    )

    # OPTION 3: Direct Quick Question (Проектная витрина)
    msg3 = (
        f"{g3}\n\n"
        f"Do you have a quick minute? My team and I built a brand-new project showcase for {clean_biz} — complete with spam-protected "
        f"quote requests, live project highlights, and seamless layout on both mobile phones and desktop PCs. "
        f"No sales pitch or strings attached, just wanted to share what we built for you.\n\n"
        f"Where is best to send the link — right here or by email? — George"
    )

    # OPTION 4: Equipment & Property Protection (Витрина техники и безопасности)
    msg4 = (
        f"{g4}\n\n"
        f"Built a dedicated, high-performance showcase specifically for {clean_biz} serving {loc_city}. "
        f"We focused heavily on eliminating homeowner fears: showcasing your {fleet}, controlled rigging over rooflines, "
        f"and guaranteed spotless turf cleanups. The showcase runs lightning-fast across all phones, tablets, and desktop monitors.\n\n"
        f"Would you be open to taking a quick look? — George"
    )

    return msg1, msg2, msg3, msg4

def get_natural_fleet_phrase(fleet, city):
    """Map internal fleet spec to an authentic arborist praise phrase."""
    fl = fleet.lower()
    loc = f"in {city}" if city and city != "Florida" else "in the area"
    
    if "crane" in fl:
        return f"crane takedowns {loc}"
    elif "spider" in fl or "compact" in fl:
        return f"tight-access lift takedowns {loc}"
    elif "emergency" in fl or "timber" in fl or "storm" in fl:
        return f"heavy timber takedowns {loc}"
    elif "stump" in fl or "grind" in fl:
        return f"takedowns and deep stump grinding {loc}"
    elif "palm" in fl or "bucket" in fl:
        return f"high-canopy takedowns {loc}"
    elif "sectional" in fl or "lowering" in fl or "rigging" in fl:
        return f"sectional takedowns {loc}"
    else:
        return f"hazardous takedowns {loc}"

def build_two_step_messages(raw_name, city, owner_name, fleet, page_url):
    """
    Build the Two-Step Outreach dialogues (Option 1: Role verification, Option 2: Fleet compliment).
    """
    clean_biz = clean_biz_name(raw_name)
    has_live_url = bool(page_url and page_url.startswith("http"))
    url_text = page_url if has_live_url else "—"

    # ==========================
    # DIALOGUE 1: ROLE VERIFICATION
    # ==========================
    if owner_name:
        d1_s1 = f"Hey {owner_name}, are you still running the crew at {clean_biz}? — George"
        if has_live_url:
            d1_s2 = (
                f"Hey {owner_name}, saw your crew's heavy rigging work and 5-star reviews in {city}. "
                f"Put together a custom showcase for {clean_biz} so homeowners can see your heavy equipment and book without haggling: "
                f"{url_text}. Wanted to share it with you directly. Take a look when you grab lunch!"
            )
        else:
            d1_s2 = (
                f"Hey {owner_name}, saw your crew's heavy rigging work and 5-star reviews in {city}. "
                f"Put together a custom showcase concept for {clean_biz} so homeowners can see your equipment and book without haggling. "
                f"Zero sales pitch, wanted to share it with you directly. Let me know if you want me to text over the preview link!"
            )
    else:
        d1_s1 = f"Hey {clean_biz} team, are you guys still taking tree removal jobs in {city}? — George"
        if has_live_url:
            d1_s2 = (
                f"Hey team, saw your crew's heavy rigging work and 5-star reviews in {city}. "
                f"Put together a custom showcase for {clean_biz} so homeowners can see your heavy equipment and book without haggling: "
                f"{url_text}. Wanted to share it with you directly. Take a look when you grab lunch!"
            )
        else:
            d1_s2 = (
                f"Hey team, saw your crew's heavy rigging work and 5-star reviews in {city}. "
                f"Put together a custom showcase concept for {clean_biz} so homeowners can see your equipment and book without haggling. "
                f"Zero sales pitch, wanted to share it with you directly. Let me know if you want me to text over the preview link!"
            )

    # ==========================
    # DIALOGUE 2: FLEET COMPLIMENT
    # ==========================
    fleet_phrase = get_natural_fleet_phrase(fleet, city)
    target = owner_name if owner_name else f"{clean_biz} team"
    d2_s1 = f"Hey {target}, saw your {fleet_phrase} — serious respect on the clean rigging. Quick question for you if you have 30 seconds? — George"

    if has_live_url:
        d2_s2 = (
            f"We actually built a dedicated project showcase for {clean_biz} to help you guys close bigger residential jobs: "
            f"{url_text}. Zero sales pitch, just wanted to get your eyes on it. Let me know what you think!"
        )
    else:
        d2_s2 = (
            f"We actually built a dedicated project showcase concept for {clean_biz} to help you guys close bigger residential jobs. "
            f"Zero sales pitch, just wanted to get your eyes on it. Let me know if you want me to text over the preview link!"
        )

    return d1_s1, d1_s2, d2_s1, d2_s2

def main():
    print("=" * 70)
    print("UPDATING USER CUSTOM TABLE: companies tree removers 250 - companies")
    print("=" * 70)

    site_configs = load_site_configs()
    leads_data = load_leads_data()
    print(f"Loaded {len(site_configs)} siteConfigs, {len(leads_data)} leads directories.")

    input_csv = os.path.join(TREE_LEADS_DIR, TARGET_CSV_NAME)
    if not os.path.exists(input_csv):
        print(f"ERROR: Input file not found: {input_csv}")
        return

    with open(input_csv, mode='r', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        header = next(reader)
        rows = list(reader)

    if header and header[0].startswith('\ufeff'):
        header[0] = header[0].replace('\ufeff', '')

    col_map = {col_name: idx for idx, col_name in enumerate(header)}
    print(f"Read {len(rows)} rows with {len(header)} columns.")

    # Base metadata columns (preserved exactly)
    base_columns = [
        'Название', 'Ссылка на Google Maps', 'Оценка', 'Количество отзывов',
        'Наличие сайта', 'Ссылка на сайт', 'Ссылка на созданную страницу',
        'Проблема с сайтом', 'Контактный номер', 'Наличие iMessage',
        'Оператор связи', 'MMS Gateway', 'Адрес', 'Штат'
    ]

    # Single-step messages (Showcase terminology, NO STOP opt-out)
    single_step_columns = [
        'Сообщение для рассылки',
        'Сообщение 1 (Готовая витрина)',
        'Сообщение 2 (Конверсия и смета)',
        'Сообщение 3 (Прямой вопрос)',
        'Сообщение 4 (Техника и безопасность)'
    ]

    # Two-step dialogues (Option 1: Role verification, Option 2: Fleet compliment)
    two_step_columns = [
        'Диалог 1: Шаг 1 (Вопрос владельцу)',
        'Диалог 1: Шаг 2 (Ответ со ссылкой)',
        'Диалог 2: Шаг 1 (Респект за технику)',
        'Диалог 2: Шаг 2 (Ответ со ссылкой)'
    ]

    final_header = base_columns + single_step_columns + two_step_columns

    processed_rows = []
    for r in rows:
        cname = r[col_map['Название']].strip() if 'Название' in col_map and col_map['Название'] < len(r) else ''
        rating = r[col_map['Оценка']].strip() if 'Оценка' in col_map and col_map['Оценка'] < len(r) else ''
        addr = r[col_map['Адрес']].strip() if 'Адрес' in col_map and col_map['Адрес'] < len(r) else ''

        ck = clean_key(cname)
        cfg = site_configs.get(ck)
        lead = leads_data.get(ck)
        folder = find_company_folder(TREE_LEADS_DIR, cname)

        owner = extract_owner(cname, lead, cfg)
        city = extract_city(cname, addr, lead, cfg)
        proof, fleet, benefit = analyze_reviews_and_equipment(folder, cfg)

        # URL to created page
        page_url = resolve_page_url(cname, lead, cfg)
        # If already present in table, keep it if valid
        if 'Ссылка на созданную страницу' in col_map and col_map['Ссылка на созданную страницу'] < len(r):
            curr_url = r[col_map['Ссылка на созданную страницу']].strip()
            if curr_url and curr_url.startswith('http'):
                page_url = curr_url

        # Build original message without STOP and using showcase
        orig_msg = build_original_message(cname, rating, city, owner, proof, fleet, benefit)

        # Enhance owner and city if inferred from original message
        if not owner and orig_msg:
            m_own = re.match(r'^Hey\s+([A-Za-z]+),', orig_msg)
            if m_own and m_own.group(1).lower() not in ['team', 'there', 'all']:
                owner = m_own.group(1)

        if (not city or city == 'Florida') and orig_msg:
            m_ct = re.search(r'(?:in|across)\s+([A-Za-z\s]+?)(?:\s+and\b|,)', orig_msg)
            if m_ct:
                city = m_ct.group(1).strip()

        # Build 4 single-step options
        m1, m2, m3, m4 = build_4_messages(cname, rating, city, owner, fleet)

        # Build two-step dialogues
        d1_s1, d1_s2, d2_s1, d2_s2 = build_two_step_messages(cname, city, owner, fleet, page_url)

        # Assemble row dict
        row_dict = {}
        for col in header:
            if col in col_map and col_map[col] < len(r):
                row_dict[col] = r[col_map[col]]

        # Override/fill message columns
        row_dict['Ссылка на созданную страницу'] = page_url
        row_dict['Сообщение для рассылки'] = orig_msg
        row_dict['Сообщение 1 (Готовая витрина)'] = m1
        row_dict['Сообщение 2 (Конверсия и смета)'] = m2
        row_dict['Сообщение 3 (Прямой вопрос)'] = m3
        row_dict['Сообщение 4 (Техника и безопасность)'] = m4
        row_dict['Диалог 1: Шаг 1 (Вопрос владельцу)'] = d1_s1
        row_dict['Диалог 1: Шаг 2 (Ответ со ссылкой)'] = d1_s2
        row_dict['Диалог 2: Шаг 1 (Респект за технику)'] = d2_s1
        row_dict['Диалог 2: Шаг 2 (Ответ со ссылкой)'] = d2_s2

        final_row = [row_dict.get(col, '') for col in final_header]
        processed_rows.append(final_row)

    # 1. Save original target CSV (overwrite the exact file user specified)
    with open(input_csv, mode='w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(final_header)
        writer.writerows(processed_rows)
    print(f"✓ Saved updated CSV: {input_csv}")

    # 2. Save clean copy without double extension
    clean_csv_path = os.path.join(TREE_LEADS_DIR, CLEAN_CSV_NAME)
    with open(clean_csv_path, mode='w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(final_header)
        writer.writerows(processed_rows)
    print(f"✓ Saved clean CSV copy: {clean_csv_path}")

    # 3. Save beautiful Excel workbook
    xlsx_path = os.path.join(TREE_LEADS_DIR, TARGET_XLSX_NAME)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tree Leads Outreach"
    ws.views.sheetView[0].showGridLines = True

    # Styling palettes
    header_fill_base = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")     # Deep Blue (base metadata)
    header_fill_single = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")   # Emerald Green (single-step showcase)
    header_fill_d1 = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")       # Indigo (Dialogue 1)
    header_fill_d2 = PatternFill(start_color="581C87", end_color="581C87", fill_type="solid")       # Purple (Dialogue 2)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    data_font = Font(name="Calibri", size=10)

    # Write headers
    for col_idx, col_name in enumerate(final_header, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        if col_name in base_columns:
            cell.fill = header_fill_base
        elif col_name in single_step_columns:
            cell.fill = header_fill_single
        elif "Диалог 1" in col_name:
            cell.fill = header_fill_d1
        else:
            cell.fill = header_fill_d2
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 36

    # Write data rows
    for r_idx, row_data in enumerate(processed_rows, 2):
        ws.row_dimensions[r_idx].height = 42
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            col_name = final_header[c_idx - 1]

            if col_name in ['Оценка', 'Количество отзывов', 'Наличие сайта', 'Наличие iMessage', 'Штат']:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_name in ['Ссылка на Google Maps', 'Ссылка на сайт', 'Ссылка на созданную страницу']:
                cell.alignment = Alignment(horizontal="left", vertical="top")
                if str(val).startswith("http"):
                    cell.font = Font(name="Calibri", size=10, color="2563EB", underline="single")
            elif "Сообщение" in col_name or "Диалог" in col_name:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top")

    # Set column widths
    width_map = {
        'Название': 26,
        'Ссылка на Google Maps': 22,
        'Оценка': 9,
        'Количество отзывов': 12,
        'Наличие сайта': 14,
        'Ссылка на сайт': 22,
        'Ссылка на созданную страницу': 30,
        'Проблема с сайтом': 22,
        'Контактный номер': 18,
        'Наличие iMessage': 15,
        'Оператор связи': 16,
        'MMS Gateway': 24,
        'Адрес': 28,
        'Штат': 8,
        'Сообщение для рассылки': 50,
        'Сообщение 1 (Готовая витрина)': 50,
        'Сообщение 2 (Конверсия и смета)': 50,
        'Сообщение 3 (Прямой вопрос)': 50,
        'Сообщение 4 (Техника и безопасность)': 50,
        'Диалог 1: Шаг 1 (Вопрос владельцу)': 40,
        'Диалог 1: Шаг 2 (Ответ со ссылкой)': 55,
        'Диалог 2: Шаг 1 (Респект за технику)': 45,
        'Диалог 2: Шаг 2 (Ответ со ссылкой)': 55
    }

    for col_idx, col_name in enumerate(final_header, 1):
        col_letter = get_column_letter(col_idx)
        w = width_map.get(col_name, 22)
        ws.column_dimensions[col_letter].width = w

    ws.freeze_panes = 'B2'

    wb.save(xlsx_path)
    print(f"✓ Saved updated Excel: {xlsx_path}")
    print("=" * 70)
    print("ALL TABLES SUCCESSFULLY UPDATED!")
    print("=" * 70)

if __name__ == "__main__":
    main()
